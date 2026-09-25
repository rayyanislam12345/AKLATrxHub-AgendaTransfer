"""Read agendas from a folder and update a local copy of the hub workbook (openpyxl)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

from . import style
from .agenda_parser import Agenda, parse_agenda
from .config import Config
from .hub import LOG_HEADERS, HubTable, UpdatePlan


def load_agendas_from_folder(folder: str | Path, on_date: date | None = None) -> list[Agenda]:
    agendas = []
    for path in sorted(Path(folder).rglob("*.docx")):
        if path.name.startswith("~$"):  # Word lock files
            continue
        agenda = parse_agenda(path)
        if on_date and agenda.agenda_date and agenda.agenda_date != on_date:
            continue
        agendas.append(agenda)
    return agendas


def read_table(ws) -> HubTable:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return HubTable(rows=rows, first_row=1, first_col=1)


def pick_sheet(wb, name: str):
    if name:
        return wb[name]
    return wb.worksheets[0]


def _side(color: str, style_name: str = "thin") -> Side:
    return Side(style=style_name, color="FF" + color)


def _style_header(cell) -> None:
    cell.font = Font(name=style.FONT, size=style.SIZE, bold=True, color="FF" + style.HEADER_FONT)
    cell.fill = PatternFill("solid", fgColor="FF" + style.HEADER_FILL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=_side(style.HEADER_DIVIDER), right=_side(style.HEADER_DIVIDER),
                         top=_side(style.HEADER_TOP, "medium"))


def _style_data(cell, align: str) -> None:
    grey = _side(style.DATA_BORDER)
    cell.font = Font(name=style.FONT, size=style.SIZE)
    cell.fill = PatternFill(fill_type=None)
    cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
    cell.border = Border(left=grey, right=grey, top=grey, bottom=grey)


def _drop_conditional_formats(ws, plan: UpdatePlan) -> None:
    """Keep other columns' conditional formatting (e.g. STATUS colours) off our columns."""
    ours = {get_column_letter(c.column) for c in plan.columns}
    old = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for cf in old:
        keep = [str(r) for r in cf.sqref.ranges
                if not ({get_column_letter(c) for c in range(range_boundaries(str(r))[0],
                                                             range_boundaries(str(r))[2] + 1)} & ours)]
        for rng in keep:
            for rule in cf.rules:
                ws.conditional_formatting.add(rng, rule)


def apply_plan(wb, ws, plan: UpdatePlan, log_sheet: str, cfg: Config | None = None) -> None:
    cfg = cfg or Config()
    _drop_conditional_formats(ws, plan)
    for col in plan.columns:
        width, align = style.layout_for(col.header, cfg)
        _style_header(ws.cell(row=col.header_row, column=col.column, value=col.header))
        ws.column_dimensions[get_column_letter(col.column)].width = width
        for offset, value in enumerate(col.values, start=1):
            cell = ws.cell(row=col.header_row + offset, column=col.column, value=value)
            if col.header_row + offset <= plan.last_row:
                _style_data(cell, align)

    if log_sheet:
        if log_sheet in wb.sheetnames:
            del wb[log_sheet]
        log = wb.create_sheet(log_sheet)
        log.append(LOG_HEADERS)
        for row in plan.log_rows:
            log.append(row)
        for idx, cell in enumerate(log[1], start=1):
            _style_header(cell)
            log.column_dimensions[get_column_letter(idx)].width = style.LOG_WIDTH
        for row in log.iter_rows(min_row=2):
            for cell in row:
                _style_data(cell, "left")


def _extend_table_or_filter(ws, plan: UpdatePlan) -> None:
    """Make the new columns filterable.

    If the header row belongs to an Excel Table, widen the table to take in the new
    columns (a second sheet filter overlapping a table corrupts the file). Otherwise
    widen, or add, the sheet's filter.
    """
    if not plan.columns:
        return
    last_col = max(c.column for c in plan.columns)
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_row != plan.header_row:
            continue
        existing = {c.name for c in table.tableColumns}
        next_id = max((c.id for c in table.tableColumns), default=0) + 1
        for col in sorted(plan.columns, key=lambda c: c.column):
            if col.column > max_col and col.header not in existing:
                table.tableColumns.append(TableColumn(id=next_id, name=col.header))
                next_id += 1
        new_max = max(max_col, last_col)
        table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(new_max)}{max_row}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
        return
    last_row = plan.header_row + len(plan.columns[0].values)
    min_col = 1
    if ws.auto_filter.ref:
        min_col, _, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        last_col, last_row = max(last_col, max_col), max(last_row, max_row)
    ws.auto_filter.ref = (f"{get_column_letter(min_col)}{plan.header_row}:"
                          f"{get_column_letter(last_col)}{last_row}")


def update_workbook(hub_path: str | Path, out_path: str | Path, build, sheet_name: str,
                    log_sheet: str, cfg: Config | None = None) -> UpdatePlan:
    """`build` is a callable HubTable -> UpdatePlan."""
    keep_vba = str(hub_path).lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(hub_path, keep_vba=keep_vba)
    ws = pick_sheet(wb, sheet_name)
    plan = build(read_table(ws))
    apply_plan(wb, ws, plan, log_sheet, cfg)
    _extend_table_or_filter(ws, plan)
    wb.save(out_path)
    return plan
