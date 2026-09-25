from datetime import date

import docx
import openpyxl
import pytest

from agenda_bot.agenda_parser import parse_agenda
from agenda_bot.config import Config
from agenda_bot.graph_backend import col_letter, col_number, share_id
from agenda_bot.hub import HubTable, build_update
from agenda_bot.local_backend import load_agendas_from_folder, update_workbook
from agenda_bot.matching import transaction_score


def make_agenda(path, lines, table_rows=None):
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=2)
        for r, (a, b) in enumerate(table_rows):
            table.cell(r, 0).text, table.cell(r, 1).text = a, b
    doc.save(path)
    return path


AGENDA = [
    "Daily Agenda", "Date: September 02, 2026", "Senior Counsel: Jane Doe", "Deliverables",
    "*Note: To be discussed with Boss.\t",
    "Transaction: ALPHA.", "Scope of Task: Amendments to the term sheet.",
    "Complete Deliverable: Revised term sheet.", "Day Deliverable: Revised term sheet.",
    "Current Deliverable Status: Pending with me.",
    "Transaction: BETA.", "Scope of Task: Finalization of the Project Documents.",
    "Complete Deliverable: Finalized Project Documents.",
    "Day Deliverable: Finalized Project Documents.", "Current Deliverable Status: Pending with me.",
    "[Subject To Allocation Of Tasks During The Day]",
    "Daily Timesheet", "Date: September 01, 2026", "Transaction: GAMMA.",
]


def test_parse_sample_layout(tmp_path):
    agenda = parse_agenda(make_agenda(tmp_path / "a.docx", AGENDA))
    assert agenda.employee == "Jane Doe"
    assert agenda.agenda_date == date(2026, 9, 2)
    assert [i.transaction for i in agenda.items] == ["ALPHA", "BETA"]  # timesheet ignored
    assert agenda.items[0].day_deliverable == "Revised term sheet"
    assert agenda.items[1].status == "Pending with me"


def test_parse_value_on_next_line_and_tables(tmp_path):
    path = make_agenda(tmp_path / "Daily Agenda - John Roe September 03 2026.docx",
                       ["Daily Agenda", "Transaction:", "ALPHA", "Transaction: Internal"],
                       table_rows=[("Transaction:", "BETA"), ("Day Deliverable:", "Memo")])
    agenda = parse_agenda(path)
    assert agenda.employee == "John Roe"                 # from file name
    assert agenda.agenda_date == date(2026, 9, 3)
    assert [i.transaction for i in agenda.items] == ["ALPHA", "BETA"]  # Internal dropped
    assert agenda.items[1].day_deliverable == "Memo"


@pytest.mark.parametrize("agenda,hub,ok", [
    ("M6", "M-6 Motorway (Hyderabad-Sukkur)", True),
    ("M6", "M60 Project", False),
    ("PTQ", "PTQ - Port Qasim", True),
    ("Karachi BRT", "BRT Karachi", False),
])
def test_transaction_score(agenda, hub, ok):
    assert (transaction_score(agenda, hub) >= 0.8) is ok


def test_alias():
    assert transaction_score("Port Qasim", "PTQ", ["Port Qasim"]) == 1.0


def hub_table():
    return HubTable(rows=[
        ["Hub title"], [],
        ["Sr", "Transaction", "Deliverable"],
        [1, "ALPHA Motorway", "Term Sheet"],
        [2, None, "Transaction Structuring Report"],
        [3, "BETA", "Project Documents"],
        [4, "BETA", "Direct Agreement"],
        [5, "DELTA", "Concession Agreement"],
        [None, None, None],
    ])


def test_build_update(tmp_path):
    a = parse_agenda(make_agenda(tmp_path / "a.docx", AGENDA))
    b = parse_agenda(make_agenda(tmp_path / "b.docx", [
        "Associate: Sam Poe", "Transaction: BETA", "Day Deliverable: Comments on project documents",
        "Current Deliverable Status: In progress", "Transaction: ALPHA",
        "Day Deliverable: Something unrelated entirely", "Transaction: OMEGA"]))
    plan = build_update(hub_table(), [a, b], Config())
    cols = {c.header: c for c in plan.columns}
    assert plan.header_row == 3
    assert [c.column for c in plan.columns] == [4, 5, 6, 7]
    working = cols["Working On Today"].values
    assert working[0] == "Jane Doe; Sam Poe"          # ALPHA term sheet + unmatched ALPHA row
    assert working[1] == "Sam Poe"                    # unmatched deliverable -> all ALPHA rows
    assert working[2] == "Jane Doe; Sam Poe"
    assert working[3] == ""
    assert cols["Deliverable Status (Today)"].values[2] == "Jane Doe: Pending with me; Sam Poe: In progress"
    assert cols["Active Today"].values == ["Yes", "Yes", "Yes", "Yes", "No", None]
    assert cols["Agenda Date"].values[0] == "2026-09-02"
    assert any(r[3] == "OMEGA" for r in plan.unmatched)


def test_existing_columns_reused():
    table = hub_table()
    table.rows[2] = ["Sr", "Transaction", "Deliverable", "Active Today"]
    table.rows[3] = [1, "ALPHA Motorway", "Term Sheet", "Yes"]
    plan = build_update(table, [], Config())
    cols = {c.header: c.column for c in plan.columns}
    assert cols["Active Today"] == 4
    assert cols["Working On Today"] == 5


def test_local_roundtrip(tmp_path):
    folder = tmp_path / "agendas"
    folder.mkdir()
    make_agenda(folder / "a.docx", AGENDA)
    make_agenda(folder / "old.docx", ["Date: August 01, 2026", "Associate: X", "Transaction: DELTA"])
    wb = openpyxl.Workbook()
    for row in hub_table().rows:
        wb.active.append(row)
    wb.save(tmp_path / "hub.xlsx")

    agendas = load_agendas_from_folder(folder, date(2026, 9, 2))
    assert [a.employee for a in agendas] == ["Jane Doe"]
    update_workbook(tmp_path / "hub.xlsx", tmp_path / "out.xlsx",
                    lambda t: build_update(t, agendas, Config()), "", "Agenda Log")
    out = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = out.worksheets[0]
    assert ws["F3"].value == "Active Today" and ws["F8"].value == "No"
    assert ws.auto_filter.ref == "A3:G8"
    assert out["Agenda Log"].max_row == 3


def test_graph_helpers():
    assert col_letter(1) == "A" and col_letter(28) == "AB" and col_number("AB") == 28
    assert share_id("https://x.sharepoint.com/a") == "u!aHR0cHM6Ly94LnNoYXJlcG9pbnQuY29tL2E"


def test_title_row_not_mistaken_for_header():
    # Regression: a merged "TRANSACTIONS HUB" banner above the real header row was
    # taken as the header, so the bot overwrote the Transaction/Deliverable columns.
    table = HubTable(rows=[
        [None] * 8,
        ["For Internal Purposes Only"] + [None] * 7,
        [None] * 8,
        ["TRANSACTIONS HUB"] + [None] * 7,
        ["S. NO.", "TRANSACTION", "DELIVERABLE", "D", "E", "WITH", "COMMENTS", "STATUS"],
        [1, "ALPHA", "Term Sheet", "d", "e", "Boss", "c", "In progress"],
    ], first_row=2)
    plan = build_update(table, [], Config())
    assert plan.header_row == 6
    assert [c.column for c in plan.columns] == [9, 10, 11, 12]


def test_never_overwrites_existing_data():
    table = HubTable(rows=[
        ["Sr", "Transaction", "Deliverable"],
        [1, "ALPHA", "Term Sheet", None, "stray note"],
    ])
    plan = build_update(table, [], Config())
    assert min(c.column for c in plan.columns) == 6


def test_graph_adds_columns_to_table():
    from agenda_bot.graph_backend import GraphWorkbook
    from agenda_bot.hub import ColumnWrite, UpdatePlan

    calls = []

    class FakeClient:
        def get(self, path, **kw):
            if path.endswith("/tables?$select=name"):
                return {"value": [{"name": "TransactionsHub"}]}
            return {"address": "'Transactions Hub'!A5:H7"}

        def request(self, method, path, **kw):
            calls.append((method, path, kw.get("json")))

    wb = GraphWorkbook.__new__(GraphWorkbook)
    wb.client, wb.base, wb.headers = FakeClient(), "/wb", {}
    plan = UpdatePlan(header_row=5, log_rows=[], columns=[
        ColumnWrite("Working On Today", 9, 5, ["A", None]),
        ColumnWrite("Active Today", 10, 5, ["Yes", "No"]),
    ])
    wb.apply_plan("Transactions Hub", plan, "")
    assert [c[0] for c in calls] == ["POST", "POST"]
    assert calls[0][1] == "/wb/tables('TransactionsHub')/columns/add"
    assert calls[0][2]["values"] == [["Working On Today"], ["A"], [""]]
    assert calls[1][2]["values"] == [["Active Today"], ["Yes"], ["No"]]


def test_local_table_is_widened(tmp_path):
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["TRANSACTIONS HUB"])
    ws.append(["Sr", "Transaction", "Deliverable"])
    ws.append([1, "ALPHA", "Term Sheet"])
    ws.add_table(Table(displayName="Hub", ref="A2:C3"))
    wb.save(tmp_path / "hub.xlsx")
    update_workbook(tmp_path / "hub.xlsx", tmp_path / "out.xlsx",
                    lambda t: build_update(t, [], Config()), "", "")
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx").active
    assert ws.tables["Hub"].ref == "A2:G3"
    assert ws.auto_filter.ref is None
    assert ws["C3"].value == "Term Sheet"
